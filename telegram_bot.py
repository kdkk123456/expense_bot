import os
try:
    from dotenv import load_dotenv
    from pathlib import Path
    load_dotenv(dotenv_path=Path(__file__).parent / ".env")
except ImportError:
    pass  # Production: env vars injected by systemd
import io
import json
import asyncio
import logging
from typing import Optional

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from telegram import Update, InputFile
from telegram.constants import ChatAction
from telegram.ext import Application, CommandHandler, MessageHandler, ContextTypes, filters

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)

TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
ADK_BASE_URL = os.getenv("ADK_BASE_URL", "http://127.0.0.1:8000")
ADK_RUNNER_URL = os.getenv("ADK_RUNNER_URL", f"{ADK_BASE_URL}/run")
ADK_APP_NAME = os.getenv("ADK_APP_NAME", "expense_bot")
ADK_USER_ID_PREFIX = os.getenv("ADK_USER_ID_PREFIX", "telegram")
REQUEST_TIMEOUT = int(os.getenv("ADK_REQUEST_TIMEOUT", "90"))

SYSTEM_INSTRUCTION = """
You are a finance assistant connected to tools.

For Telegram responses:
- Never use markdown pipe tables like | col | col |.
- Prefer short answers, bullets, or <pre>...</pre> fixed-width blocks.
- Keep outputs readable on mobile.
- If the user asks for Excel, xlsx, export, spreadsheet, file, or download, use export_expenses or export_income.
- For export requests, return the tool output directly as raw JSON only.
"""

def is_export_request(text: str) -> bool:
    lowered = text.lower()
    keywords = ["excel", "xlsx", "export", "download", "spreadsheet", "file"]
    return any(word in lowered for word in keywords)

def guess_period(text: str) -> str:
    lowered = text.lower()
    if "last month" in lowered:
        return "last_month"
    if "this year" in lowered or "yearly" in lowered or "current year" in lowered:
        return "this_year"
    if "last year" in lowered:
        return "last_year"
    return "this_month"

def normalize_export_prompt(text: str) -> str:
    period = guess_period(text)
    if "income" in text.lower():
        return f"Export income for {period} as Excel."
    return f"Export expenses for {period} as Excel."

async def send_typing_indicator(chat_id: int, bot, stop_event: asyncio.Event):
    while not stop_event.is_set():
        try:
            await bot.send_chat_action(chat_id=chat_id, action=ChatAction.TYPING)
        except Exception:
            pass
        try:
            await asyncio.wait_for(stop_event.wait(), timeout=4)
        except asyncio.TimeoutError:
            continue

def extract_reply_text(events) -> str:
    if not isinstance(events, list):
        return ""

    collected_texts = []

    for event in events:
        if not isinstance(event, dict):
            continue
        content = event.get("content") or {}
        parts = content.get("parts") or []

        if isinstance(parts, list):
            for part in parts:
                if isinstance(part, dict):
                    text = part.get("text")
                    if isinstance(text, str) and text.strip():
                        collected_texts.append(text.strip())

    for text in reversed(collected_texts):
        if "<pre>" in text and "</pre>" in text:
            return text

    return collected_texts[-1] if collected_texts else ""

def extract_export_payload(events) -> Optional[dict]:
    if not isinstance(events, list):
        return None

    for event in reversed(events):
        if not isinstance(event, dict):
            continue
        content = event.get("content") or {}
        parts = content.get("parts") or []

        if not isinstance(parts, list):
            continue

        for part in parts:
            if not isinstance(part, dict):
                continue
            text = part.get("text")
            if not isinstance(text, str):
                continue

            try:
                data = json.loads(text)
            except Exception:
                continue

            if (
                isinstance(data, dict)
                and data.get("export_type") in {"expenses", "income"}
                and isinstance(data.get("rows"), list)
            ):
                return data

    return None

def autosize_columns(ws):
    for column_cells in ws.columns:
        cells = list(column_cells)
        if not cells:
            continue

        max_length = 0
        column_letter = cells[0].column_letter

        for cell in cells:
            try:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))
            except Exception:
                pass

        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 30)

def build_excel_file(export_payload: dict) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("Sheet1")

    export_type = export_payload.get("export_type", "export")
    period = export_payload.get("period", "data")
    rows = export_payload.get("rows", [])

    if export_type == "expenses":
        ws.title = "Expenses"
        headers = [
            "Expense ID",
            "Description",
            "Category",
            "Amount",
            "Vendor",
            "Date",
            "Status",
            "Notes",
            "Created At",
        ]
        ws.append(headers)

        for row in rows:
            ws.append([
                row.get("expense_id"),
                row.get("description"),
                row.get("category"),
                row.get("amount"),
                row.get("vendor_paid_to"),
                row.get("date"),
                row.get("status"),
                row.get("notes"),
                str(row.get("created_at") or ""),
            ])

        filename = f"expenses_{period}.xlsx"
    else:
        ws.title = "Income"
        headers = [
            "Income ID",
            "Customer Name",
            "Category",
            "Amount",
            "Description",
            "Date",
            "Status",
            "Notes",
            "Created At",
        ]
        ws.append(headers)

        for row in rows:
            ws.append([
                row.get("income_id"),
                row.get("customer_name"),
                row.get("category"),
                row.get("amount"),
                row.get("description"),
                row.get("date"),
                row.get("status"),
                row.get("notes"),
                str(row.get("created_at") or ""),
            ])

        filename = f"income_{period}.xlsx"

    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = Font(bold=True)

    autosize_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    output.name = filename
    return output

def ensure_adk_session(user_id: str):
    session_url = f"{ADK_BASE_URL}/apps/{ADK_APP_NAME}/users/{user_id}/sessions/{user_id}"
    response = requests.post(session_url, json={}, timeout=REQUEST_TIMEOUT)

    if response.status_code not in (200, 201, 409):
        print("SESSION STATUS:", response.status_code)
        print("SESSION RESPONSE:", response.text)
        response.raise_for_status()

def call_adk(user_id: str, message_text: str):
    ensure_adk_session(user_id)

    payload = {
        "appName": ADK_APP_NAME,
        "userId": user_id,
        "sessionId": user_id,
        "newMessage": {
            "role": "user",
            "parts": [
                {
                    "text": f"{SYSTEM_INSTRUCTION.strip()}\n\nUser request: {message_text}"
                }
            ],
        },
    }

    response = requests.post(
        ADK_RUNNER_URL,
        json=payload,
        timeout=REQUEST_TIMEOUT,
    )

    if not response.ok:
        print("ADK STATUS:", response.status_code)
        print("ADK RESPONSE:", response.text)

    response.raise_for_status()
    return response.json()

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    message = update.message or update.effective_message
    if message is None:
        return

    await message.reply_text(
        "Send a finance request like:\n"
        "- show expenses this month\n"
        "- summary this year\n"
        "- export expenses this month to excel"
    )

async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    message = update.effective_message
    user = update.effective_user

    if message is None or user is None:
        return

    text = (message.text or "").strip()
    if not text:
        await message.reply_text("Please send a text message.")
        return

    user_id = f"{ADK_USER_ID_PREFIX}-{user.id}"

    stop_event = asyncio.Event()
    typing_task = asyncio.create_task(
        send_typing_indicator(message.chat_id, context.bot, stop_event)
    )

    try:
        effective_prompt = normalize_export_prompt(text) if is_export_request(text) else text
        events = await asyncio.to_thread(call_adk, user_id, effective_prompt)

        print("ADK EVENTS START")
        print(json.dumps(events, indent=2, ensure_ascii=False))
        print("ADK EVENTS END")

        export_payload = extract_export_payload(events)
        if export_payload:
            file_buffer = build_excel_file(export_payload)
            stop_event.set()
            await typing_task
            await message.reply_document(
                document=InputFile(file_buffer, filename=file_buffer.name),
                caption=f"Here is your {export_payload['export_type']} export.",
            )
            return

        reply_text = extract_reply_text(events) or "I could not generate a response."
        stop_event.set()
        await typing_task

        try:
            await message.reply_text(reply_text, parse_mode="HTML")
        except Exception:
            await message.reply_text(reply_text)

    except requests.HTTPError as e:
        logger.exception("ADK HTTP error")
        stop_event.set()
        await typing_task
        await message.reply_text(f"Server error: {e}")
    except Exception as e:
        logger.exception("Unhandled error")
        stop_event.set()
        await typing_task
        await message.reply_text(f"Something went wrong: {e}")

def main():
    if not TELEGRAM_BOT_TOKEN:
        raise RuntimeError("Missing TELEGRAM_BOT_TOKEN")

    app = Application.builder().token(TELEGRAM_BOT_TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))
    app.run_polling()

if __name__ == "__main__":
    main()